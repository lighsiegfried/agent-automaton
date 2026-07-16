"""Per-format extraction to structured blocks (Phase 6A, items 2-3).

TXT/Markdown/CSV are handled with the stdlib; PDF (pypdf) and DOCX (python-docx) are
lazy imports so a host without them degrades gracefully to a clear failure rather than
a crash. Extraction preserves page (PDF), heading/section (Markdown/DOCX), and row
(CSV) metadata so chunks can be cited precisely. OCR is never invoked automatically.
"""

from __future__ import annotations

import csv
import io
import re

from app.core import errors
from app.knowledge.models import (
    FORMAT_CSV, FORMAT_DOCX, FORMAT_MD, FORMAT_PDF, FORMAT_TXT, Block, ExtractedDoc,
)


class ExtractionError(Exception):
    def __init__(self, code: str, reason: str):
        super().__init__(reason)
        self.code = code
        self.reason = reason


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _title_from(filename: str) -> str:
    base = (filename or "document").rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    return base.rsplit(".", 1)[0] or base


# --- TXT ---------------------------------------------------------------------------

def _extract_txt(data: bytes, filename: str) -> ExtractedDoc:
    text = _decode(data)
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    blocks = [Block(text=p) for p in paras] or [Block(text=text.strip())]
    return ExtractedDoc(title=_title_from(filename), blocks=blocks, page_count=1)


# --- Markdown ----------------------------------------------------------------------

def _extract_md(data: bytes, filename: str) -> ExtractedDoc:
    text = _decode(data)
    blocks: list[Block] = []
    section = None
    buf: list[str] = []

    def flush():
        if buf:
            blocks.append(Block(text="\n".join(buf).strip(), section=section, heading=section))
            buf.clear()

    title = _title_from(filename)
    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            section = m.group(2).strip()
            if not title or title == _title_from(filename):
                title = title if blocks else section
        else:
            buf.append(line)
    flush()
    blocks = [b for b in blocks if b.text]
    return ExtractedDoc(title=title or _title_from(filename), blocks=blocks or [Block(text=text.strip())],
                        page_count=1)


# --- CSV ---------------------------------------------------------------------------

def _extract_csv(data: bytes, filename: str) -> ExtractedDoc:
    text = _decode(data)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise ExtractionError(errors.EMPTY_DOCUMENT, "the CSV has no rows")
    header = rows[0]
    blocks: list[Block] = []
    for i, row in enumerate(rows[1:], start=1):
        pairs = "; ".join(f"{(header[j] if j < len(header) else f'col{j}')}: {val}"
                          for j, val in enumerate(row) if val)
        if pairs:
            blocks.append(Block(text=pairs, row=i))
    if not blocks:
        raise ExtractionError(errors.EMPTY_DOCUMENT, "the CSV has no data rows")
    return ExtractedDoc(title=_title_from(filename), blocks=blocks, page_count=1)


# --- PDF ---------------------------------------------------------------------------

def _extract_pdf(data: bytes, filename: str, max_pages: int) -> ExtractedDoc:
    try:
        from pypdf import PdfReader
    except Exception:
        raise ExtractionError(errors.EXTRACTION_FAILED, "PDF support (pypdf) is not installed")
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(errors.EXTRACTION_FAILED, f"could not read the PDF: {type(exc).__name__}")
    if getattr(reader, "is_encrypted", False):
        raise ExtractionError(errors.ENCRYPTED_FILE, "the PDF is encrypted")
    pages = reader.pages
    if len(pages) > max_pages:
        raise ExtractionError(errors.TOO_MANY_PAGES, f"{len(pages)} pages exceeds {max_pages}")
    blocks: list[Block] = []
    for i, page in enumerate(pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        if text:
            blocks.append(Block(text=text, page=i))
    if not blocks:
        raise ExtractionError(errors.EMPTY_DOCUMENT,
                              "no extractable text (a scanned PDF needs OCR, which is disabled)")
    return ExtractedDoc(title=_title_from(filename), blocks=blocks, page_count=len(pages))


# --- DOCX --------------------------------------------------------------------------

def _extract_docx(data: bytes, filename: str) -> ExtractedDoc:
    try:
        import docx
    except Exception:
        raise ExtractionError(errors.EXTRACTION_FAILED, "DOCX support (python-docx) is not installed")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(errors.EXTRACTION_FAILED, f"could not read the DOCX: {type(exc).__name__}")
    blocks: list[Block] = []
    section = None
    title = None
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name if para.style else "") or ""
        if style.startswith("Heading") or style == "Title":
            section = text
            if title is None:
                title = text
            continue
        blocks.append(Block(text=text, section=section, heading=section))
    for t, table in enumerate(document.tables, start=1):
        for r, row in enumerate(table.rows, start=1):
            cells = "; ".join(c.text.strip() for c in row.cells if c.text.strip())
            if cells:
                blocks.append(Block(text=cells, section=section or f"table {t}", row=r))
    if not blocks:
        raise ExtractionError(errors.EMPTY_DOCUMENT, "the DOCX has no extractable text")
    return ExtractedDoc(title=title or _title_from(filename), blocks=blocks, page_count=1)


def extract(data: bytes, fmt: str, filename: str, *, max_pages: int = 1000) -> ExtractedDoc:
    if fmt == FORMAT_TXT:
        return _extract_txt(data, filename)
    if fmt == FORMAT_MD:
        return _extract_md(data, filename)
    if fmt == FORMAT_CSV:
        return _extract_csv(data, filename)
    if fmt == FORMAT_PDF:
        return _extract_pdf(data, filename, max_pages)
    if fmt == FORMAT_DOCX:
        return _extract_docx(data, filename)
    raise ExtractionError(errors.UNSUPPORTED_FORMAT, f"no extractor for {fmt}")
