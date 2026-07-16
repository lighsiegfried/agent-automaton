"""End-to-end Knowledge Vault (Phase 6A): the API (upload/search/ask/list/delete),
PDF/DOCX page-and-section citations, NL routing, and Activity Center integration."""

import io
import itertools
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from app.activity import events as activity_events
from app.activity import service as activity_service
from app.activity.service import ActivityService
from app.config import get_settings
from app.core import pending
from app.core.router import handle_command
from app.knowledge import service as ksvc_mod
from app.knowledge.embeddings import HashEmbedder
from app.knowledge.service import KnowledgeService
from app.main import app
from app.schemas.commands import CommandRequest, Intent

client = TestClient(app)
UTC = timezone.utc


def make_pdf(text: str) -> bytes:
    esc = lambda s: s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = ("BT /F1 12 Tf 72 720 Td 14 TL\n"
              + "\n".join(f"({esc(l)}) Tj T*" for l in text.split("\n")) + "\nET").encode("latin-1")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = b"%PDF-1.4\n"
    offs = []
    for i, body in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xp = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for o in offs:
        out += f"{o:010d} 00000 n \n".encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\nstartxref\n"
            + str(xp).encode() + b"\n%%EOF")
    return out


def make_docx(heading: str, paras: list[str]) -> bytes:
    import docx
    d = docx.Document()
    d.add_heading(heading, level=1)
    for p in paras:
        d.add_paragraph(p)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


@pytest.fixture
def env(tmp_path, monkeypatch):
    s = get_settings()
    monkeypatch.setattr(s, "enable_knowledge_vault", True)
    monkeypatch.setattr(s, "knowledge_owner", "local")
    monkeypatch.setattr(s, "knowledge_db_path", tmp_path / "k.db")
    monkeypatch.setattr(s, "knowledge_files_dir", tmp_path / "files")
    monkeypatch.setattr(s, "knowledge_indexes_dir", tmp_path / "idx")
    monkeypatch.setattr(s, "enable_activity_center", True)
    monkeypatch.setattr(s, "activity_owner", "local")
    monkeypatch.setattr(s, "activity_db_path", tmp_path / "a.db")
    monkeypatch.setattr(pending, "_broker", pending.PendingBroker())
    ids = (f"{i:032x}" for i in itertools.count(1))
    ksvc = KnowledgeService(db_path=tmp_path / "k.db", files_dir=tmp_path / "files",
                            indexes_dir=tmp_path / "idx", embedder=HashEmbedder(),
                            now_fn=lambda: datetime(2026, 7, 15, tzinfo=UTC), id_factory=lambda: next(ids))
    monkeypatch.setattr(ksvc_mod, "_service", ksvc)
    asvc = ActivityService(db_path=tmp_path / "a.db",
                           now_fn=lambda: datetime(2026, 7, 15, 12, tzinfo=UTC), settings_provider=lambda: s)
    monkeypatch.setattr(activity_service, "_service", asvc)
    activity_events.install()
    return SimpleNamespace(settings=s, ksvc=ksvc, asvc=asvc)


def _upload(filename, data, collections=""):
    return client.post("/knowledge/documents/upload",
                       files={"file": (filename, data)}, data={"collections": collections}).json()


# --- API upload + all formats ------------------------------------------------------


def test_upload_txt_md_csv_pdf_docx(env):
    assert _upload("notes.txt", b"BOFASA client. Payment net 30.")["status"] == "ok"
    assert _upload("policy.md", b"# Security\n\nPasswords rotate every 90 days.")["status"] == "ok"
    assert _upload("data.csv", b"name,role\nAna,lead")["status"] == "ok"
    pdf = _upload("report.pdf", make_pdf("BOFASA revenue grew twelve percent.\nThe audit passed."))
    assert pdf["status"] == "ok" and pdf["data"]["document"]["page_count"] == 1
    docx = _upload("sec.docx", make_docx("Retention Policy", ["Retention is 30 days.", "Backups run nightly."]))
    assert docx["status"] == "ok"
    assert client.get("/knowledge/documents").json()["data"]["count"] == 5


def test_pdf_answer_cites_page(env):
    _upload("report.pdf", make_pdf("The security audit passed in Q3.\nRevenue grew twelve percent."))
    ans = client.post("/knowledge/ask", json={"question": "did the security audit pass"}).json()
    assert ans["data"]["grounded"] is True
    assert any("page" in c["citation"] for c in ans["data"]["citations"])


def test_docx_cites_section(env):
    _upload("sec.docx", make_docx("Data Retention", ["The retention period is thirty days for logs."]))
    res = client.post("/knowledge/search", json={"query": "retention period logs"}).json()
    assert res["data"]["evidence"] and res["data"]["evidence"][0]["section"] == "Data Retention"


def test_reject_before_upload(env):
    r = _upload("malware.exe", b"MZ\x90\x00")
    assert r["error_code"] == "UNSUPPORTED_FORMAT"


# --- NL routing --------------------------------------------------------------------


def test_nl_routing(env):
    _upload("notes.txt", b"BOFASA client. Payment net 30.")
    assert handle_command(CommandRequest(text="busca en mis documentos información sobre BOFASA",
                                         language="es")).intent == Intent.KNOWLEDGE_SEARCH
    assert handle_command(CommandRequest(text="resume la política de seguridad",
                                         language="es")).intent == Intent.KNOWLEDGE_ASK
    assert handle_command(CommandRequest(text="lista mis documentos",
                                         language="es")).intent == Intent.KNOWLEDGE_LIST


# --- delete through API ------------------------------------------------------------


def test_delete_prepare_confirm_via_api(env):
    did = _upload("a.txt", b"BOFASA client. Payment net 30.")["data"]["document"]["doc_id"]
    assert client.request("DELETE", f"/knowledge/documents/{did}/prepare").json()["status"] == "needs_confirmation"
    done = client.request("DELETE", f"/knowledge/documents/{did}/confirm",
                          json={"phrase": "confirmar eliminación de documento"}).json()
    assert done["status"] == "executed"
    assert client.get(f"/knowledge/documents/{did}").json()["document"] is None


# --- Activity Center integration ---------------------------------------------------


def test_ingestion_and_search_appear_in_activity(env):
    _upload("notes.txt", b"BOFASA client. Payment net 30.")
    client.post("/knowledge/search", json={"query": "BOFASA"})
    events = env.asvc.events(page=1, page_size=50)["events"]
    kinds = {e["event_type"] for e in events if e["domain"] == "knowledge"}
    assert "document_ready" in kinds and "knowledge_search" in kinds


def test_status_reports_health(env):
    _upload("notes.txt", b"BOFASA client.")
    st = client.get("/knowledge/status").json()["data"]
    assert st["enabled"] is True and st["counts"]["ready"] == 1
    assert st["ocr_enabled"] is False and st["index_health"] == "ok"
